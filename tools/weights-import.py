#!/usr/bin/env python3
"""Перенос веса из таблицы владельца в справочник моделей.

Таблица — docs/weights/Ива-модели-вес.xlsx (или .csv с теми же колонками): «Бренд», «Модель», «Вес, кг», «kind (не менять)».
Вес пишется в supabase/models/[1-5]-*.txt как weight=… (кг). Пустая ячейка — вес не трогаем.
После переноса: python3 tools/models-gen.py — и SQL готов.

  python3 tools/weights-import.py ФАЙЛ            перенести
  python3 tools/weights-import.py ФАЙЛ --dry      только показать, что изменится
"""
import csv, glob, os, re, sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
COLS = {"brand": "Бренд", "model": "Модель", "weight": "Вес, кг", "kind": "kind (не менять)"}
MIN_KG, MAX_KG = 0.05, 20000


def num(v):
    """«2,9» / «2.9» / 2.9 → «2.9»; пусто → None; не число → ValueError."""
    if v is None:
        return None
    s = str(v).strip().replace(",", ".").replace(" ", "").replace("\u00a0", "")
    s = re.sub(r"(?i)кг$", "", s)
    if not s:
        return None
    x = float(s)
    if not (MIN_KG <= x <= MAX_KG):
        raise ValueError("вес вне 0,05–20000 кг")
    t = ("%.3f" % x).rstrip("0").rstrip(".")
    return t


def read_rows(path):
    if path.lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        it = ws.iter_rows(values_only=True)
        head = [str(h or "").strip() for h in next(it)]
        for r in it:
            yield dict(zip(head, r))
    else:
        with open(path, encoding="utf-8-sig", newline="") as f:
            sample = f.read(4096); f.seek(0)
            dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
            for r in csv.DictReader(f, dialect=dialect):
                yield {k.strip(): v for k, v in r.items() if k}


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    dry = "--dry" in sys.argv
    if len(args) != 1:
        print(__doc__); return 2
    want, errors = {}, []
    for i, r in enumerate(read_rows(args[0]), 2):
        kind, brand, model = (str(r.get(COLS[k]) or "").strip() for k in ("kind", "brand", "model"))
        if not kind or not brand or not model:
            continue
        try:
            w = num(r.get(COLS["weight"]))
        except ValueError as e:
            errors.append("строка %d (%s %s): %s — «%s»" % (i, brand, model, e, r.get(COLS["weight"])))
            continue
        if w is not None:
            want[(kind, brand, model)] = w
    changed, same, found = 0, 0, set()
    for f in sorted(glob.glob(os.path.join(ROOT, "supabase/models/[1-5]-*.txt"))):
        lines = open(f, encoding="utf-8").read().split("\n")
        dirty = False
        for n, l in enumerate(lines):
            if l.startswith("#") or "|" not in l:
                continue
            parts = l.split("|")
            key = tuple(p.strip() for p in parts[:3])
            if key not in want:
                continue
            found.add(key)
            specs = [s.strip() for s in parts[3].split(";") if s.strip()]
            old = next((s.split("=", 1)[1] for s in specs if s.startswith("weight=")), None)
            if old == want[key]:
                same += 1
                continue
            specs = [s for s in specs if not s.startswith("weight=")] + ["weight=" + want[key]]
            parts[3] = " " + "; ".join(specs) + (" " if len(parts) > 4 else "")
            lines[n] = "|".join(parts)
            print("%s %s %s: %s → %s" % (key[0], key[1], key[2], old or "—", want[key]))
            changed += 1; dirty = True
        if dirty and not dry:
            open(f, "w", encoding="utf-8").write("\n".join(lines))
    missing = sorted(set(want) - found)
    for k in missing:
        errors.append("нет в справочнике: %s | %s | %s" % k)
    for e in errors:
        print("✘ " + e)
    print("ИТОГ: с весом в таблице %d, изменено %d, без изменений %d, ошибок %d%s"
          % (len(want), changed, same, len(errors), " (проверка, файлы не тронуты)" if dry else ""))
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
